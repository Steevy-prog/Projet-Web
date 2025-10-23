from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_color_cell(table, row_idx, col_idx, hex_color, color_name, usage):
    """Ajoute une cellule colorée au tableau"""
    cell = table.rows[row_idx].cells[col_idx]
    
    # Convertir hex en RGB
    hex_color = hex_color.lstrip('#')
    r, g, b = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    
    # Appliquer la couleur de fond
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), hex_color.upper())
    cell._element.get_or_add_tcPr().append(shading_elm)
    
    # Ajouter le texte
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(f"{color_name}\n{hex_color.upper()}\n{usage}")
    run.font.size = Pt(9)
    
    # Couleur du texte (blanc ou noir selon la luminosité)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    if luminance > 0.5:
        run.font.color.rgb = RGBColor(0, 0, 0)
    else:
        run.font.color.rgb = RGBColor(255, 255, 255)

# Créer le document
doc = Document()

# Titre principal
title = doc.add_heading('Charte Graphique', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(212, 166, 71)

subtitle = doc.add_paragraph('Restaurant Élégance - Application Web')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.runs[0]
subtitle_run.font.size = Pt(14)
subtitle_run.font.color.rgb = RGBColor(160, 160, 160)

doc.add_paragraph()

# 1. IDENTITÉ VISUELLE
doc.add_heading('1. Identité Visuelle', level=1)

doc.add_paragraph(
    'La charte graphique de Restaurant Élégance repose sur un design minimaliste et sophistiqué, '
    'alliant élégance et modernité. Le thème sombre avec des accents dorés crée une atmosphère '
    'raffinée et exclusive, évoquant le luxe et la qualité.'
)

# 2. PALETTE DE COULEURS
doc.add_heading('2. Palette de Couleurs Principale', level=1)

doc.add_heading('2.1 Couleurs de Base', level=2)

# Tableau des couleurs principales
table = doc.add_table(rows=6, cols=3)
table.style = 'Light Grid Accent 1'

# En-têtes
headers = table.rows[0].cells
headers[0].text = 'Aperçu'
headers[1].text = 'Couleur'
headers[2].text = 'Utilisation'

# Couleurs principales
colors_main = [
    ('#0b0b0d', 'Background', 'Fond principal de l\'application'),
    ('#151518', 'Card', 'Cartes et conteneurs'),
    ('#d4a647', 'Primary', 'Couleur primaire - Or élégant'),
    ('#f5f5f5', 'Foreground', 'Texte principal'),
    ('#a0a0a0', 'Muted', 'Texte secondaire et désactivé')
]

for idx, (hex_col, name, usage) in enumerate(colors_main, 1):
    add_color_cell(table, idx, 0, hex_col, name, '')
    table.rows[idx].cells[1].text = f'{name}\n{hex_col.upper()}'
    table.rows[idx].cells[2].text = usage

doc.add_paragraph()

doc.add_heading('2.2 Palette Dorée (Variations)', level=2)

doc.add_paragraph('Gamme complète des tons dorés utilisés pour créer de la profondeur :')

gold_palette = [
    ('#b88b1f', 'Or Ancien', 'Tons sombres, éléments subtils'),
    ('#d4a647', 'Or Principal', 'Couleur primaire, boutons, liens'),
    ('#f4c563', 'Or Accent', 'Highlights, états hover'),
    ('#ffd97d', 'Or Clair', 'Accents lumineux, effets de brillance'),
    ('#8b6914', 'Or Profond', 'Ombres, graphiques')
]

for hex_col, name, usage in gold_palette:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} ({hex_col.upper()}) - ')
    run.bold = True
    p.add_run(usage)

doc.add_paragraph()

doc.add_heading('2.3 Couleurs Fonctionnelles', level=2)

functional_colors = [
    ('#1a1a1d', 'Secondary', 'Arrière-plans secondaires'),
    ('#d4183d', 'Destructive', 'Erreurs, suppressions, alertes'),
    ('rgba(212,166,71,0.35)', 'Border', 'Bordures des éléments'),
    ('rgba(244,197,99,0.65)', 'Ring', 'Focus, anneaux de sélection')
]

for hex_col, name, usage in functional_colors:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} ({hex_col}) - ')
    run.bold = True
    p.add_run(usage)

# 3. TYPOGRAPHIE
doc.add_heading('3. Typographie', level=1)

doc.add_heading('3.1 Polices', level=2)

fonts = [
    ('Sans-serif système', 'ui-sans-serif, system-ui, sans-serif', 'Police principale pour tout le texte'),
    ('Monospace', 'ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas', 'Code et données techniques')
]

for font_name, font_stack, usage in fonts:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{font_name} - ')
    run.bold = True
    p.add_run(f'{usage}\n  Stack: {font_stack}')

doc.add_heading('3.2 Échelle Typographique', level=2)

type_scale = [
    ('6xl', '3.75rem (60px)', 'Titres héros, pages d\'accueil'),
    ('5xl', '3rem (48px)', 'Titres principaux'),
    ('4xl', '2.25rem (36px)', 'Sous-titres importants'),
    ('3xl', '1.875rem (30px)', 'Titres de sections'),
    ('2xl', '1.5rem (24px)', 'Titres H1'),
    ('xl', '1.25rem (20px)', 'Titres H2'),
    ('lg', '1.125rem (18px)', 'Titres H3'),
    ('base', '1rem (16px)', 'Texte standard'),
    ('sm', '0.875rem (14px)', 'Texte secondaire'),
    ('xs', '0.75rem (12px)', 'Petits textes, labels')
]

table_typo = doc.add_table(rows=len(type_scale)+1, cols=3)
table_typo.style = 'Light List Accent 1'

headers = table_typo.rows[0].cells
headers[0].text = 'Taille'
headers[1].text = 'Valeur'
headers[2].text = 'Usage'

for idx, (size, value, usage) in enumerate(type_scale, 1):
    table_typo.rows[idx].cells[0].text = size
    table_typo.rows[idx].cells[1].text = value
    table_typo.rows[idx].cells[2].text = usage

doc.add_paragraph()

doc.add_heading('3.3 Poids de Police', level=2)

weights = [
    ('Normal (400)', 'Texte courant, paragraphes'),
    ('Medium (500)', 'Titres, boutons, labels, emphase')
]

for weight, usage in weights:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{weight} - ')
    run.bold = True
    p.add_run(usage)

# 4. ESPACEMENTS ET GRILLE
doc.add_heading('4. Espacements et Grille', level=1)

doc.add_heading('4.1 Système d\'Espacement', level=2)

doc.add_paragraph('Basé sur une unité de base de 0.25rem (4px) :')

spacing = [
    ('1', '0.25rem (4px)', 'Espacement minimal'),
    ('2', '0.5rem (8px)', 'Espacement petit'),
    ('4', '1rem (16px)', 'Espacement standard'),
    ('6', '1.5rem (24px)', 'Espacement moyen'),
    ('8', '2rem (32px)', 'Espacement large'),
    ('12', '3rem (48px)', 'Espacement très large'),
    ('16', '4rem (64px)', 'Espacement extra large')
]

for unit, value, usage in spacing:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'Unité {unit} ({value}) - ')
    run.bold = True
    p.add_run(usage)

doc.add_heading('4.2 Conteneurs', level=2)

containers = [
    ('md', '28rem (448px)', 'Petits conteneurs'),
    ('2xl', '42rem (672px)', 'Conteneurs moyens'),
    ('3xl', '48rem (768px)', 'Conteneurs standards'),
    ('4xl', '56rem (896px)', 'Grands conteneurs'),
    ('6xl', '72rem (1152px)', 'Conteneurs très larges')
]

for size, value, usage in containers:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{size} ({value}) - ')
    run.bold = True
    p.add_run(usage)

# 5. BORDURES ET ARRONDIS
doc.add_heading('5. Bordures et Arrondis', level=1)

doc.add_heading('5.1 Radius (Coins Arrondis)', level=2)

doc.add_paragraph('Valeur de base : 1rem (16px)')

radius = [
    ('sm', 'calc(1rem - 4px) = 12px', 'Petits éléments'),
    ('md', 'calc(1rem - 2px) = 14px', 'Éléments moyens'),
    ('lg', '1rem = 16px', 'Éléments standards'),
    ('xl', 'calc(1rem + 4px) = 20px', 'Grands éléments'),
    ('2xl', '1rem = 16px', 'Cartes principales')
]

for size, value, usage in radius:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{size} ({value}) - ')
    run.bold = True
    p.add_run(usage)

doc.add_heading('5.2 Bordures', level=2)

borders = [
    ('Border standard', '2px solid rgba(212, 166, 71, 0.35)', 'Bordures par défaut'),
    ('Border hover', '2px solid rgba(244, 197, 99, 0.65)', 'État hover'),
    ('Border gradient', 'Gradient doré (135deg)', 'Éléments premium')
]

for name, value, usage in borders:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} - ')
    run.bold = True
    p.add_run(f'{value} - {usage}')

# 6. OMBRES ET EFFETS
doc.add_heading('6. Ombres et Effets', level=1)

doc.add_heading('6.1 Box Shadows', level=2)

shadows = [
    ('Shadow standard', '0 0 #0000', 'Pas d\'ombre par défaut'),
    ('Shadow hover', '0 10px 25px rgba(212,166,71,0.3)', 'Élévation au hover'),
    ('Shadow gold pulse', '0 0 20px rgba(244,197,99,0.5)', 'Animation pulsation'),
    ('Shadow card hover', '0 12px 30px rgba(212,166,71,0.25)', 'Cartes au hover')
]

for name, value, usage in shadows:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} - ')
    run.bold = True
    p.add_run(f'{value}\n  Usage: {usage}')

doc.add_heading('6.2 Effets de Flou', level=2)

blurs = [
    ('sm', '8px', 'Flou léger'),
    ('lg', '16px', 'Flou prononcé')
]

for size, value, usage in blurs:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'Blur {size} ({value}) - ')
    run.bold = True
    p.add_run(usage)

# 7. ANIMATIONS
doc.add_heading('7. Animations et Transitions', level=1)

doc.add_heading('7.1 Timing et Durées', level=2)

timing = [
    ('Durée par défaut', '0.15s', 'Transitions rapides'),
    ('Timing function', 'cubic-bezier(0.4, 0, 0.2, 1)', 'Courbe d\'accélération naturelle'),
    ('Animations longues', '2-3s', 'Animations décoratives (pulse, shine)')
]

for name, value, usage in timing:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} - ')
    run.bold = True
    p.add_run(f'{value} - {usage}')

doc.add_heading('7.2 Animations Personnalisées', level=2)

animations = [
    ('goldPulse', '2s ease-in-out infinite', 'Pulsation dorée pour éléments importants'),
    ('goldShine', '3s ease-in-out infinite', 'Effet de brillance qui traverse l\'élément'),
    ('goldSparkle', '2s ease-in-out infinite', 'Scintillement subtil'),
    ('goldGlow', '2s ease-in-out infinite', 'Lueur dorée sur le texte'),
    ('fadeInScale', '0.5s ease-out', 'Apparition avec zoom'),
    ('slideUp', '0.6s ease-out', 'Glissement vers le haut'),
    ('borderGlow', '2s ease-in-out infinite', 'Bordure qui pulse'),
    ('shimmer', '2s infinite', 'Effet de chargement')
]

for name, duration, description in animations:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{name} ({duration}) - ')
    run.bold = True
    p.add_run(description)

doc.add_heading('7.3 Effets Hover', level=2)

hovers = [
    ('hover-gold-lift', 'translateY(-4px) + shadow', 'Élévation avec ombre'),
    ('hover-gold-scale', 'scale(1.05)', 'Agrandissement'),
    ('hover-gold-glow', 'box-shadow glow', 'Lueur dorée'),
    ('hover-gold-brighten', 'brightness(1.2)', 'Éclaircissement'),
    ('Boutons', 'translateY(-2px)', 'Légère élévation'),
    ('Cartes', 'translateY(-6px) + shadow', 'Élévation prononcée')
]

for effect, transform, description in hovers:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{effect} - ')
    run.bold = True
    p.add_run(f'{transform} - {description}')

# 8. COMPOSANTS UI
doc.add_heading('8. Composants UI', level=1)

doc.add_heading('8.1 Boutons', level=2)

buttons = [
    ('Primary', 'Fond doré (#d4a647), texte sombre', 'Actions principales'),
    ('Secondary', 'Fond sombre (#1a1a1d), texte clair', 'Actions secondaires'),
    ('Destructive', 'Fond rouge (#d4183d), texte blanc', 'Actions destructives'),
    ('Ghost', 'Transparent, texte doré', 'Actions tertiaires'),
    ('Outline', 'Bordure dorée, fond transparent', 'Actions alternatives')
]

for btn_type, style, usage in buttons:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{btn_type} - ')
    run.bold = True
    p.add_run(f'{style}\n  Usage: {usage}')

doc.add_heading('8.2 Cartes', level=2)

cards = [
    ('Background', '#151518 (card)', 'Fond des cartes'),
    ('Border', 'rgba(212,166,71,0.35)', 'Bordure subtile'),
    ('Radius', '1rem (16px)', 'Coins arrondis'),
    ('Padding', '1.5rem (24px)', 'Espacement interne'),
    ('Hover', 'translateY(-6px) + shadow', 'Animation au survol')
]

for prop, value, description in cards:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{prop} - ')
    run.bold = True
    p.add_run(f'{value} - {description}')

doc.add_heading('8.3 Formulaires', level=2)

forms = [
    ('Input background', '#1a1a1d', 'Fond des champs'),
    ('Input border', 'rgba(212,166,71,0.45)', 'Bordure des champs'),
    ('Focus ring', 'rgba(244,197,99,0.65)', 'Anneau de focus'),
    ('Label', 'Medium (500), base (16px)', 'Style des labels'),
    ('Placeholder', '#a0a0a0', 'Texte placeholder')
]

for element, value, description in forms:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{element} - ')
    run.bold = True
    p.add_run(f'{value} - {description}')

# 9. ICONOGRAPHIE
doc.add_heading('9. Iconographie', level=1)

doc.add_paragraph(
    'Bibliothèque : Lucide React v0.487.0\n'
    'Style : Outline (contours)\n'
    'Taille par défaut : 24px\n'
    'Couleur : Hérite du texte parent ou doré (#d4a647)'
)

icon_usage = [
    ('Navigation', 'Menu, Home, User, Settings, LogOut'),
    ('Actions', 'Plus, Trash, Edit, Check, X'),
    ('Commerce', 'ShoppingCart, CreditCard, Package'),
    ('Social', 'Instagram, Facebook, Twitter, Github'),
    ('Stats', 'TrendingUp, BarChart, PieChart, Activity'),
    ('Jeux', 'Gamepad2, Trophy, Star, Gift')
]

for category, icons in icon_usage:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{category} - ')
    run.bold = True
    p.add_run(icons)

# 10. GRILLES ET LAYOUTS
doc.add_heading('10. Grilles et Layouts', level=1)

doc.add_heading('10.1 Responsive Breakpoints', level=2)

breakpoints = [
    ('sm', '640px', 'Petits mobiles'),
    ('md', '768px', 'Tablettes'),
    ('lg', '1024px', 'Petits écrans'),
    ('xl', '1280px', 'Écrans standards'),
    ('2xl', '1536px', 'Grands écrans')
]

for bp, value, device in breakpoints:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{bp} ({value}) - ')
    run.bold = True
    p.add_run(device)

doc.add_heading('10.2 Grilles', level=2)

grids = [
    ('Grid 1 colonne', 'Mobile par défaut'),
    ('Grid 2 colonnes', 'md: et plus'),
    ('Grid 3 colonnes', 'lg: et plus'),
    ('Grid 4 colonnes', 'xl: et plus (menus, jeux)'),
    ('Gap standard', '1.5rem (24px)')
]

for grid, usage in grids:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{grid} - ')
    run.bold = True
    p.add_run(usage)

# 11. ACCESSIBILITÉ
doc.add_heading('11. Accessibilité', level=1)

accessibility = [
    ('Contrastes', 'Conformité WCAG 2.1 AA minimum'),
    ('Navigation clavier', 'Tous les éléments interactifs accessibles au clavier'),
    ('ARIA labels', 'Labels descriptifs sur tous les composants'),
    ('Focus visible', 'Anneau doré visible sur focus'),
    ('Tailles tactiles', 'Minimum 44x44px pour les boutons'),
    ('Texte alternatif', 'Images avec attribut alt descriptif')
]

for criterion, description in accessibility:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{criterion} - ')
    run.bold = True
    p.add_run(description)

# 12. PRINCIPES DE DESIGN
doc.add_heading('12. Principes de Design', level=1)

principles = [
    ('Minimalisme', 'Grands espaces négatifs, design épuré, focus sur l\'essentiel'),
    ('Élégance', 'Thème sombre sophistiqué avec accents dorés raffinés'),
    ('Cohérence', 'Utilisation systématique des composants et styles'),
    ('Hiérarchie visuelle', 'Tailles, poids et couleurs pour guider l\'attention'),
    ('Feedback visuel', 'Animations et transitions pour confirmer les actions'),
    ('Responsive', 'Adaptation fluide à tous les formats d\'écran'),
    ('Performance', 'Animations optimisées, chargement rapide'),
    ('Accessibilité', 'Design inclusif pour tous les utilisateurs')
]

for principle, description in principles:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(f'{principle} - ')
    run.bold = True
    p.add_run(description)

# 13. EXEMPLES D'UTILISATION
doc.add_heading('13. Exemples d\'Utilisation', level=1)

doc.add_heading('13.1 Carte de Menu', level=2)

menu_card = [
    'Background: #151518 (card)',
    'Border: 2px solid rgba(212,166,71,0.35)',
    'Radius: 1rem',
    'Padding: 1.5rem',
    'Image: Aspect ratio 16:9, radius 0.75rem',
    'Titre: text-xl (20px), medium (500), couleur foreground',
    'Prix: text-2xl (24px), medium (500), couleur primary (#d4a647)',
    'Hover: translateY(-6px) + shadow doré'
]

for spec in menu_card:
    doc.add_paragraph(spec, style='List Bullet')

doc.add_heading('13.2 Bouton Principal', level=2)

primary_btn = [
    'Background: #d4a647 (primary)',
    'Texte: #0b0b0d (primary-foreground), medium (500)',
    'Padding: 0.5rem 1rem',
    'Radius: 0.5rem',
    'Hover: translateY(-2px), brightness(1.1)',
    'Active: translateY(0)',
    'Transition: 0.3s cubic-bezier(0.4, 0, 0.2, 1)'
]

for spec in primary_btn:
    doc.add_paragraph(spec, style='List Bullet')

# CONCLUSION
doc.add_heading('14. Conclusion', level=1)

doc.add_paragraph(
    'Cette charte graphique définit l\'identité visuelle de Restaurant Élégance, '
    'créant une expérience utilisateur cohérente, élégante et accessible. '
    'Le thème sombre avec accents dorés évoque le luxe et la sophistication, '
    'tout en restant moderne et fonctionnel.'
)

doc.add_paragraph()

doc.add_paragraph(
    'Tous les éléments de cette charte doivent être respectés pour maintenir '
    'la cohérence visuelle de l\'application. Les animations et micro-interactions '
    'ajoutent une touche de raffinement sans compromettre les performances.'
)

# Sauvegarder
doc.save('Charte_Graphique_Restaurant_Elegance.docx')
print('✅ Charte graphique créée : Charte_Graphique_Restaurant_Elegance.docx')
